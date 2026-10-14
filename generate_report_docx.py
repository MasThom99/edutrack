import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_placeholder_box(doc, caption_text, box_height_inches=2.2):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(5.8)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:left w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:right w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run(f"[ TEMPAT TANGKAPAN LAYAR (SCREENSHOT) ]\n{caption_text}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Caption paragraph below table
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(9.5)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def build_laporan_docx():
    target_path = r"C:\Users\thom\.gemini\antigravity\scratch\edutrack\Laporan_Praktikum_Pertemuan_1_Ikhwan_Nur_Rizki_Fathama.docx"
    doc = docx.Document()

    # Page Margins (Format Standar Akademik Indonesia: Kiri 4cm, Atas 4cm, Kanan 3cm, Bawah 3cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.57)     # ~4 cm
        section.bottom_margin = Inches(1.18)  # ~3 cm
        section.left_margin = Inches(1.57)    # ~4 cm
        section.right_margin = Inches(1.18)   # ~3 cm

    # Styling Defaults
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # ==========================================
    # HALAMAN COVER
    # ==========================================
    p_cov_title = doc.add_paragraph()
    p_cov_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_title.paragraph_format.space_before = Pt(20)
    p_cov_title.paragraph_format.space_after = Pt(6)
    r1 = p_cov_title.add_run("LAPORAN PRAKTIKUM\nPEMROGRAMAN MOBILE LANJUT")
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_cov_sub = doc.add_paragraph()
    p_cov_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_sub.paragraph_format.space_before = Pt(6)
    p_cov_sub.paragraph_format.space_after = Pt(24)
    r2 = p_cov_sub.add_run("PERTEMUAN 1: REVIEW PMO DASAR & PRETEST\nIMPLEMENTASI RELASI DATA ONE-TO-MANY PADA APLIKASI EDUTRACK\n(STUDENT ACADEMIC & SKILL PORTFOLIO TRACKER)")
    r2.font.bold = True
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Placeholder Logo Kampus
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(30)
    p_logo.paragraph_format.space_after = Pt(40)
    r_logo = p_logo.add_run("[ TEMPAT LOGO UNIVERSITAS AMIKOM PURWOKERTO ]")
    r_logo.font.italic = True
    r_logo.font.size = Pt(11)
    r_logo.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Data Penyusun
    p_disusun = doc.add_paragraph()
    p_disusun.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disusun.paragraph_format.space_before = Pt(20)
    p_disusun.paragraph_format.space_after = Pt(4)
    r_dis = p_disusun.add_run("Disusun Oleh:")
    r_dis.font.size = Pt(11)

    p_mhs = doc.add_paragraph()
    p_mhs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mhs.paragraph_format.space_before = Pt(2)
    p_mhs.paragraph_format.space_after = Pt(20)
    r_mhs = p_mhs.add_run("IKHWAN NUR RIZKI FATHAMA\nNIM: 24SA11A023")
    r_mhs.font.bold = True
    r_mhs.font.size = Pt(12)

    p_dos = doc.add_paragraph()
    p_dos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dos.paragraph_format.space_before = Pt(4)
    p_dos.paragraph_format.space_after = Pt(36)
    r_dos = p_dos.add_run("Dosen Pengampu:\nNandang Hermanto, M.Kom.")
    r_dos.font.size = Pt(11)
    r_dos.font.bold = True

    p_kampus = doc.add_paragraph()
    p_kampus.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kampus.paragraph_format.space_before = Pt(24)
    p_kampus.paragraph_format.space_after = Pt(0)
    r_kmp = p_kampus.add_run("PROGRAM STUDI INFORMATIKA\nFAKULTAS ILMU KOMPUTER\nUNIVERSITAS AMIKOM PURWOKERTO\nPURWOKERTO\n2025/2026")
    r_kmp.font.bold = True
    r_kmp.font.size = Pt(12)
    r_kmp.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_page_break()

    # ==========================================
    # BAB I: PENDAHULUAN
    # ==========================================
    def add_chapter_header(title, subtitle=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        if subtitle:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(16)
            run2 = p2.add_run(subtitle)
            run2.font.bold = True
            run2.font.size = Pt(13)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.bold = True
        run.font.italic = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    def add_body(text, indent=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11.5)

    add_chapter_header("BAB I", "PENDAHULUAN")
    
    add_h2("1.1 Latar Belakang")
    add_body(
        "Perkembangan teknologi perangkat bergerak (mobile devices) telah mengubah paradigma pengelolaan dan konsumsi informasi di berbagai sektor, termasuk lingkungan akademik perguruan tinggi. Mahasiswa kini dituntut tidak hanya berfokus pada capaian indeks prestasi kumulatif (IPK) semata, melainkan juga rekam jejak sertifikasi keahlian, kompetensi industri, dan pencapaian portofolio pendukung yang terverifikasi."
    )
    add_body(
        "Mata kuliah Pemrograman Mobile Lanjut pada Pertemuan 1 menyelenggarakan latihan praktikum sekaligus pretest untuk menguji kesiapan mahasiswa dalam merancang dan membangun arsitektur aplikasi mobile secara utuh. Fokus utama penugasan ini adalah menginterpretasikan kebutuhan sistem dari deskripsi yang abstrak, merancang struktur relasi data bertingkat (One-to-Many), serta mengimplementasikan antarmuka interaktif yang memenuhi standar kaidah rekayasa perangkat lunak mobile."
    )
    add_body(
        "Untuk menjawab tantangan tersebut, dibangunlah aplikasi bernama 'EduTrack' (Student Academic & Skill Portfolio Tracker). EduTrack dirancang untuk memfasilitasi penelusuran profil direktori mahasiswa sekaligus merinci rekam jejak kompetensi sertifikasi keahlian yang dimiliki masing-masing mahasiswa secara lokal, terstruktur, dan elegan menggunakan framework Flutter."
    )

    add_h2("1.2 Tujuan Praktikum")
    add_body(
        "Tujuan pelaksanaan praktikum ini mengacu secara langsung pada target capaian pembelajaran modul pertemuan 1, antara lain:", indent=False
    )
    add_body("1. Menunjukkan pemahaman dasar terhadap pengembangan aplikasi mobile berbasis Flutter dan bahasa pemrograman Dart.", indent=False)
    add_body("2. Mampu menganalisis dan menginterpretasikan kebutuhan sistem dari deskripsi abstrak menjadi model data konkret.", indent=False)
    add_body("3. Mampu mendesain struktur data bertingkat dengan relasi satu-ke-banyak (One-to-Many) antara entitas utama mahasiswa dan entitas sertifikasi.", indent=False)
    add_body("4. Mampu merancang alur navigasi antar layar yang efisien dengan teknik passing object.", indent=False)
    add_body("5. Mengimplementasikan komponen visual modern berbasis Material 3 yang menyajikan informasi secara ringkas, estetis, dan informatif.", indent=False)

    add_h2("1.3 Batasan Masalah")
    add_body(
        "Agar pengerjaan tugas tetap terarah dan memenuhi kriteria teknis modul praktikum, batasan masalah yang ditetapkan adalah:", indent=False
    )
    add_body("1. Aplikasi dibangun menggunakan framework Flutter dengan arsitektur kode modular berorientasi objek (OOP).", indent=False)
    add_body("2. Data yang digunakan merupakan data lokal statis (mock/dummy data) tanpa ketergantungan koneksi jaringan backend (API) maupun database eksternal.", indent=False)
    add_body("3. Entitas utama memuat data identitas mahasiswa (NIM, Nama, Jurusan, Email, Foto Profil via Network Image).", indent=False)
    add_body("4. Entitas relasi memuat sekurang-kurangnya 5 data item sertifikasi/kompetensi yang merepresentasikan pencapaian untuk setiap mahasiswa.", indent=False)
    add_body("5. Alur navigasi terdiri dari dua layar utama: Layar 1 (Daftar Mahasiswa) dan Layar 2 (Detail Profil & Portofolio Sertifikasi).", indent=False)

    doc.add_page_break()

    # ==========================================
    # BAB II: LANDASAN TEORI
    # ==========================================
    add_chapter_header("BAB II", "LANDASAN TEORI")

    add_h2("2.1 Framework Flutter dan Bahasa Pemrograman Dart")
    add_body(
        "Flutter merupakan Software Development Kit (SDK) antarmuka pengguna sumber terbuka (open-source) yang dikembangkan oleh Google untuk membangun aplikasi multiplatform (Android, iOS, Web, Desktop) dari satu basis kode tunggal (single codebase). Flutter menggunakan bahasa pemrograman Dart yang mendukung pemrograman berorientasi objek murni, pengetikan statis yang kuat (sound null safety), serta kompilasi Ahead-of-Time (AOT) untuk menghasilkan performa native yang gegas."
    )

    add_h2("2.2 Konsep Declarative UI dan Widget Tree")
    add_body(
        "Berbeda dengan paradigma imperatif konvensional, Flutter mengusung paradigma Declarative UI. Pada arsitektur ini, antarmuka pengguna direfleksikan sebagai fungsi langsung dari kondisi state aplikasi: UI = f(state). Seluruh elemen visual di dalam Flutter dibungkus ke dalam komponen yang disebut Widget. Pohon hierarki widget (Widget Tree) mengatur susunan tata letak dari elemen akar (root/MaterialApp) hingga elemen daun terkecil seperti Text, Icon, maupun Container."
    )

    add_h2("2.3 Pemodelan Relasi Data Satu-ke-Banyak (One-to-Many)")
    add_body(
        "Relasi data satu-ke-banyak (One-to-Many Relationship) merupakan pola pemodelan di mana satu data rekaman pada entitas utama (Parent Entity) dapat diasosiasikan dengan nol, satu, atau banyak data rekaman pada entitas lain (Child Entity). Dalam paradigma Object-Oriented Programming (OOP) pada Dart, relasi ini diimplementasikan secara elegan dengan menyematkan koleksi berulang bertipe List<Child> sebagai properti di dalam definisi class Parent tanpa membutuhkan skema tabel database relasional yang rumit."
    )

    add_h2("2.4 Navigasi Layar dan Object Parameter Passing")
    add_body(
        "Mekanisme perpindahan antar layar di Flutter dikelola oleh widget Navigator yang beroperasi menyerupai struktur data tumpukan (Stack LIFO). Untuk memindahkan pengguna ke layar baru, digunakan metode Navigator.push(context, MaterialPageRoute(...)). Keunggulan utama pendekatan OOP pada Flutter adalah kemampuan mengirimkan instance objek utuh secara langsung ke konstruktor layar tujuan (Object Parameter Passing). Hal ini menghindarkan aplikasi dari keharusan melakukan query pencarian ulang ID yang memboroskan sumber daya memori."
    )

    add_h2("2.5 Prinsip Desain Antarmuka Material 3")
    add_body(
        "Material 3 (Material You) merupakan evolusi terbaru dari sistem desain Google yang menekankan personalisasi dinamis, kontras warna yang nyaman, radius sudut kurva yang lembut (rounded corners), serta hierarki visual yang jelas. Pada aplikasi mobile modern, penerapan Material 3 meningkatkan keterbacaan (readability) data kompleks melalui penggunaan Card, Chip, Badge, dan ListTile dengan elevasi halus."
    )

    doc.add_page_break()

    # ==========================================
    # BAB III: PERANCANGAN DAN ARSITEKTUR APLIKASI
    # ==========================================
    add_chapter_header("BAB III", "PERANCANGAN DAN ARSITEKTUR APLIKASI")

    add_h2("3.1 Identitas dan Konsep Aplikasi")
    add_body(
        "Aplikasi yang dirancang bernama **EduTrack (Student Academic & Skill Portfolio Tracker)**. Aplikasi ini difungsikan sebagai platform direktori akademik modern yang menampilkan ikhtisar identitas mahasiswa serta merinci jejak rekam kompetensi dan sertifikasi keahlian yang telah diraih. Pendekatan ini mengangkat konteks nyata kebutuhan perguruan tinggi dalam memetakan kesiapan karier mahasiswa di industri."
    )

    add_h2("3.2 Perancangan Struktur Data (One-to-Many)")
    add_body(
        "Struktur data aplikasi dimodelkan secara terpisah dan bertingkat menggunakan dua kelas cetak biru (blueprint):", indent=False
    )
    add_body("1. **Entitas Utama (`Student`)**:", indent=False)
    add_body("   - `id`: Nomor Induk Mahasiswa / NIM (kunci unik bertipe String).", indent=False)
    add_body("   - `name`: Nama lengkap mahasiswa (String).", indent=False)
    add_body("   - `major`: Program studi / Jurusan (String).", indent=False)
    add_body("   - `email`: Alamat surel resmi mahasiswa (String).", indent=False)
    add_body("   - `imageUrl`: Tautan alamat foto profil daring yang dimuat melalui NetworkImage (String).", indent=False)
    add_body("   - `certifications`: Kumpulan data sertifikasi milik mahasiswa bersangkutan bertipe `List<Certification>`.", indent=False)
    add_body("2. **Entitas Relasi (`Certification`)**:", indent=False)
    add_body("   - `title`: Nama bidang keahlian atau judul sertifikasi (String).", indent=False)
    add_body("   - `issuer`: Organisasi/lembaga penerbit kredensial seperti Dicoding, Google, Cisco, BNSP (String).", indent=False)
    add_body("   - `year`: Tahun perolehan sertifikasi (String).", indent=False)
    add_body("   - `grade`: Predikat kelulusan atau hasil asesmen seperti Kompeten, Sangat Baik, Grade A (String).", indent=False)

    add_h2("3.3 Struktur Direktori Modular Proyek")
    add_body(
        "Untuk memenuhi standar Clean Code dan keterawatan (maintainability), struktur berkas proyek diatur ke dalam folder modular sebagai berikut:", indent=False
    )
    
    # Table of Project Structure
    table_struct = doc.add_table(rows=7, cols=2)
    table_struct.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_struct.autofit = False
    
    headers = ["Lokasi Berkas", "Fungsi & Peranan Modul"]
    col_widths = [Inches(2.2), Inches(3.6)]
    
    for i, h in enumerate(headers):
        cell = table_struct.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    struct_rows = [
        ("lib/models/certification_model.dart", "Mendefinisikan entitas relasi sertifikasi kompetensi."),
        ("lib/models/student_model.dart", "Mendefinisikan entitas utama mahasiswa dengan List<Certification>."),
        ("lib/data/dummy_data.dart", "Menyediakan data statis 4 mahasiswa dengan 5+ sertifikasi terstruktur."),
        ("lib/screens/student_list_screen.dart", "Layar 1: Menampilkan daftar kartu profil mahasiswa yang dapat digulir."),
        ("lib/screens/student_detail_screen.dart", "Layar 2: Menampilkan profil lengkap & portofolio sertifikasi (1-to-N)."),
        ("lib/main.dart", "Entry point aplikasi, konfigurasi tema Material 3 Academic Modern.")
    ]

    for row_idx, (path_val, desc_val) in enumerate(struct_rows, start=1):
        c1 = table_struct.cell(row_idx, 0)
        c2 = table_struct.cell(row_idx, 1)
        c1.width = col_widths[0]
        c2.width = col_widths[1]
        bg_c = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        set_cell_background(c1, bg_c)
        set_cell_background(c2, bg_c)
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c2, top=80, bottom=80, left=100, right=100)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(path_val)
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(desc_val)
        r2.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("3.4 Alur Navigasi Pengguna (User Flow)")
    add_body(
        "Alur pengoperasian aplikasi dirancang ringkas dalam dua tahap navigasi utama:", indent=False
    )
    add_body("1. **Layar 1 (Daftar Mahasiswa)**: Pengguna membuka aplikasi dan disajikan daftar kartu profil mahasiswa. Tiap kartu merangkum foto profil, nama, NIM, jurusan, serta indikator total sertifikat.", indent=False)
    add_body("2. **Perpindahan Navigasi**: Pengguna mengetuk salah satu kartu mahasiswa. Aplikasi mengeksekusi `Navigator.push` dengan meneruskan seluruh instansiasi objek `Student` terkait.", indent=False)
    add_body("3. **Layar 2 (Detail Profil & Portofolio)**: Menampilkan header profil berdimensi besar dan daftar kartu sertifikasi (minimal 5 item) lengkap dengan penerbit, tahun, dan badge grade kelulusan.", indent=False)
    add_body("4. **Navigasi Kembali**: Pengguna menekan tombol 'Back' di AppBar untuk kembali ke daftar utama tanpa kehilangan riwayat gulir.", indent=False)

    add_h2("3.5 Perancangan Antarmuka & Skema Warna")
    add_body(
        "Aplikasi menerapkan palet warna bertema **Academic Modern**:", indent=False
    )
    add_body("• **Warna Utama (Primary - #1E3A8A / Deep Indigo)**: Diterapkan pada AppBar, header profil, dan tombol identitas utama untuk memberikan impresi profesional, formal, dan stabil.", indent=False)
    add_body("• **Warna Aksen (Accent - #059669 / Emerald Green)**: Diterapkan pada badge status/grade kelulusan sertifikasi untuk melambangkan pencapaian kompetensi yang teruji.", indent=False)
    add_body("• **Warna Latar (Background - #F8FAFC / Soft Slate Neutral)**: Bidang kanvas lembut yang memberi efek kedalaman visual saat dipadukan dengan kartu-kartu putih berelevasi rendah.", indent=False)

    doc.add_page_break()

    # ==========================================
    # BAB IV: IMPLEMENTASI DAN PEMBAHASAN
    # ==========================================
    add_chapter_header("BAB IV", "IMPLEMENTASI DAN PEMBAHASAN")

    add_h2("4.1 Implementasi Kode Program dan Penjelasan Modul")
    add_body(
        "Berikut disajikan penjelasan dan dokumentasi kode sumber untuk masing-masing modul berkas yang membangun aplikasi EduTrack:"
    )

    # 1. certification_model.dart
    add_h3("A. Model Relasi Sertifikasi (lib/models/certification_model.dart)")
    add_body(
        "Kelas `Certification` berfungsi sebagai cetak biru untuk entitas data sertifikasi. Atribut dibuat bertipe final dengan konstruktor konstan untuk memastikan immutability data saat aplikasi berjalan.", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.1 Tangkapan Layar Kode certification_model.dart")

    # 2. student_model.dart
    add_h3("B. Model Entitas Mahasiswa (lib/models/student_model.dart)")
    add_body(
        "Kelas `Student` menampung informasi identitas mahasiswa serta properti `certifications` bertipe `List<Certification>`. Di sinilah implementasi relasi One-to-Many diwujudkan pada tingkat kode objek.", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.2 Tangkapan Layar Kode student_model.dart")

    # 3. dummy_data.dart
    add_h3("C. Modul Data Statis (lib/data/dummy_data.dart)")
    add_body(
        "Berkas ini menyediakan dataset lokal yang memuat 4 mahasiswa berprestasi dari beragam program studi (Teknik Informatika, Sistem Informasi, Teknik Komputer, Sains Data). Setiap mahasiswa dilengkapi minimal 5 riwayat sertifikasi kredibel.", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.3 Tangkapan Layar Kode dummy_data.dart")

    # 4. student_list_screen.dart
    add_h3("D. Layar Daftar Mahasiswa (lib/screens/student_list_screen.dart)")
    add_body(
        "Menyajikan komponen `ListView.builder` yang merender daftar kartu mahasiswa secara dinamis. Di dalamnya terdapat fungsi `onTap` yang memanggil `Navigator.push` menuju `StudentDetailScreen` dengan membawa objek mahasiswa yang dipilih.", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.4 Tangkapan Layar Kode student_list_screen.dart")

    # 5. student_detail_screen.dart
    add_h3("E. Layar Detail Profil & Portofolio (lib/screens/student_detail_screen.dart)")
    add_body(
        "Menerima parameter objek `Student`. Layar ini menampilkan header visual profil mahasiswa dan melakukan perulangan `ListView.builder` bertingkat untuk menyajikan daftar sertifikasi lengkap dengan badge predikat berwarna hijau zamrud (Emerald Green).", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.5 Tangkapan Layar Kode student_detail_screen.dart")

    # 6. main.dart
    add_h3("F. Konfigurasi Tema & Entry Point (lib/main.dart)")
    add_body(
        "Menginisialisasi konfigurasi `MaterialApp`, menetapkan `ThemeData` berbasis Material 3 dengan palet warna Deep Indigo dan Soft Slate, serta menetapkan `StudentListScreen` sebagai layar awal.", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.6 Tangkapan Layar Kode main.dart")

    add_h2("4.2 Hasil Tampilan Antarmuka dan Pengujian Sistem")
    add_body(
        "Pengujian antarmuka dilakukan secara langsung melalui web server dan peramban browser Chrome. Aplikasi berhasil merender tata letak secara responsif dan interaktif sesuai dengan desain visual yang direncanakan."
    )

    add_placeholder_box(doc, "Gambar 4.7 Hasil Antarmuka Layar 1 (Daftar Mahasiswa EduTrack)")
    add_body(
        "Pada Gambar 4.7, terlihat tampilan Layar 1 menyajikan kumpulan kartu profil mahasiswa yang bersih dan rapi. Tiap kartu menyajikan avatar lingkaran, nama mahasiswa, NIM, jurusan, serta chip penanda total sertifikasi yang dimiliki.", indent=False
    )

    add_placeholder_box(doc, "Gambar 4.8 Hasil Antarmuka Layar 2 (Detail Profil & Riwayat Sertifikasi 1-to-N)")
    add_body(
        "Pada Gambar 4.8, tampak tampilan Layar 2 saat salah satu kartu mahasiswa dipilih. Header profil menampilkan avatar berdimensi besar dan rincian kontak surel. Bagian bawah memvisualisasikan rekam jejak 5 sertifikasi kompetensi lengkap dengan badge status kelulusan berkontras tinggi.", indent=False
    )

    add_h3("Hasil Pengujian Statis (flutter analyze)")
    add_body(
        "Pengujian kualitas kode program dilakukan menggunakan perkakas analisis statis Flutter compiler bawaan dengan perintah `flutter analyze`. Hasil eksekusi menunjukkan status bersih tanpa adanya galat linter, tipe data, maupun peringatan sintaksis:", indent=False
    )
    add_placeholder_box(doc, "Gambar 4.9 Hasil Analisis Statis flutter analyze (No issues found!)")

    add_h2("4.3 Analisis Mandiri (Menjawab 4 Pertanyaan Pemicu Modul)")
    add_body(
        "Berdasarkan instruksi modul praktikum pertemuan 1, berikut adalah analisis mandiri komprehensif untuk menjawab 4 pertanyaan pemicu evaluasi sistem:"
    )

    # Question 1
    add_h3("1. Bagaimana cara Anda merancang struktur data agar fleksibel dan mudah dikembangkan?")
    add_body(
        "Struktur data dirancang menggunakan prinsip Pemrograman Berorientasi Objek (OOP) dengan memisahkan entitas utama (`Student`) dan entitas relasi (`Certification`) ke dalam berkas model independen. Model `Student` menampung koleksi objek `List<Certification>`. Pendekatan modular ini memberikan fleksibilitas tinggi: jika di masa mendatang diperlukan penambahan atribut baru (seperti tanggal kedaluwarsa sertifikat, nomor registrasi akreditasi, atau portofolio proyek terpisah), pengembang cukup menambahkan atribut pada model relasi tanpa mengganggu struktur data mahasiswa utama."
    )

    # Question 2
    add_h3("2. Bagaimana Anda memastikan navigasi antar halaman tetap efisien dan mudah dipahami pengguna?")
    add_body(
        "Efisiensi navigasi dicapai dengan menerapkan teknik *Object Parameter Passing* pada perintah `Navigator.push`. Seluruh instance objek `Student` yang telah berada di memori diteruskan langsung ke konstruktor `StudentDetailScreen`. Pendekatan ini mengeliminasi kebutuhan proses pencarian data (query lookup) berbasis ID di layar tujuan yang memakan waktu komputasi. Dari perspektif pengguna, navigasi dibuat intuitif dengan kartu profil yang dapat diklik langsung (tappable card) dan tombol kembali (Back Arrow) di AppBar kiri atas yang memelihara riwayat posisi scroll pada daftar sebelumnya."
    )

    # Question 3
    add_h3("3. Bagaimana Anda menampilkan informasi yang banyak secara ringkas namun tetap informatif?")
    add_body(
        "Penyajian informasi yang padat diatasi dengan strategi *Progressive Disclosure* dan pembedaan hierarki visual. Pada Layar 1 (Daftar Mahasiswa), informasi ditampilkan secara ringkas hanya memuat poin-poin kunci pengenal (foto, nama, NIM, jurusan, dan chip counter jumlah sertifikat) agar pengguna dapat memindai direktori dengan cepat tanpa beban kognitif berlebih. Rincian mendalam (lembaga penerbit, tahun perolehan, dan predikat nilai sertifikasi) dialokasikan seluruhnya pada Layar 2 (Detail Profil), di mana informasi dikelompokkan ke dalam kartu-kartu terstruktur berelevasi lembut."
    )

    # Question 4
    add_h3("4. Apa pendekatan terbaik untuk mengelola dan menampilkan data yang saling berkaitan?")
    add_body(
        "Pendekatan terbaik untuk relasi data satu-ke-banyak (One-to-Many) dalam konteks antarmuka mobile adalah *Nested List Composition*. Data relasi diisolasi di dalam wadah objek entitas induknya, kemudian pada antarmuka ditampilkan menggunakan komponen `ListView.builder` yang disematkan di dalam `SingleChildScrollView` (dengan parameter `shrinkWrap: true` dan `physics: NeverScrollableScrollPhysics`). Pendekatan ini menjamin keterpaduan data (data integrity), mencegah percampuran data antar mahasiswa, serta menghasilkan pengalaman pengguliran layar yang mulus dan bebas lag visual."
    )

    doc.add_page_break()

    # ==========================================
    # BAB V: PENUTUP
    # ==========================================
    add_chapter_header("BAB V", "PENUTUP")

    add_h2("5.1 Kesimpulan")
    add_body(
        "Berdasarkan perancangan, implementasi, dan pengujian yang telah dilaksanakan pada tugas praktikum dan pretest Pemrograman Mobile Lanjut Pertemuan 1, dapat ditarik beberapa kesimpulan penting:", indent=False
    )
    add_body("1. Aplikasi **EduTrack** telah berhasil dibangun menggunakan framework Flutter dan bahasa Dart dengan memenuhi seluruh spesifikasi tugas yang disyaratkan.", indent=False)
    add_body("2. Konsep relasi data satu-ke-banyak (One-to-Many) berhasil diimplementasikan secara solid melalui integrasi `List<Certification>` di dalam kelas model `Student`.", indent=False)
    add_body("3. Alur navigasi antar layar beroperasi secara efektif dan gegas memanfaatkan teknik passing object instan pada `Navigator.push`.", indent=False)
    add_body("4. Desain visual berbasis Material 3 dengan tema Academic Modern (Deep Indigo dan Emerald Green) memberikan tampilan antarmuka yang bersih, profesional, dan mudah dipahami.", indent=False)
    add_body("5. Hasil verifikasi statis melalui `flutter analyze` mengonfirmasi bahwa seluruh basis kode bebas dari galat sintaksis dan sesuai dengan standar konvensi pemrograman Flutter modern.", indent=False)

    add_h2("5.2 Saran dan Rencana Pengembangan Lanjutan")
    add_body(
        "Untuk pengembangan aplikasi EduTrack di masa mendatang, disarankan beberapa peningkatan fitur:", indent=False
    )
    add_body("1. **Integrasi Basis Data Lokal (Persistence Storage)**: Mengintegrasikan pustaka basis data lokal seperti SQLite (sqflite) atau Hive/Isar agar data mahasiswa dan sertifikasi dapat disimpan secara permanen di media penyimpanan perangkat.", indent=False)
    add_body("2. **Konektivitas Layanan Jaringan (RESTful API)**: Menghubungkan aplikasi dengan backend server berbasis REST API menggunakan protokol HTTP/Dio untuk mendukung sinkronisasi data daring secara real-time.", indent=False)
    add_body("3. **Fitur Formulir Input Data (CRUD)**: Menambahkan fitur formulir interaktif bagi pengguna untuk menambah, mengubah, dan menghapus data sertifikasi secara dinamis.", indent=False)
    add_body("4. **Fitur Pencarian dan Filter**: Menyediakan kolom pencarian (search bar) dan penyaringan kategori sertifikasi untuk mempermudah penelusuran data saat jumlah mahasiswa bertambah banyak.", indent=False)

    doc.add_page_break()

    # ==========================================
    # DAFTAR PUSTAKA
    # ==========================================
    add_chapter_header("DAFTAR PUSTAKA")
    
    pustaka_items = [
        "Flutter Documentation. (2024). <i>Flutter architectural overview & Widget Catalog</i>. Google Developers. Diakses dari https://docs.flutter.dev/",
        "Hermanto, N. (2025). <i>Modul Praktikum Pertemuan 1: Review PMO Dasar (Pemrograman Mobile Lanjut)</i>. Purwokerto: Universitas Amikom Purwokerto.",
        "Material Design Team. (2024). <i>Material 3 Design Guidelines: Colors, Components, and Layouts</i>. Google LLC. Diakses dari https://m3.material.io/",
        "Phillips, B., Stewart, C., & Marsicano, K. (2021). <i>Android Programming: The Big Nerd Ranch Guide</i>. Big Nerd Ranch Guides.",
        "Smyth, N. (2023). <i>Flutter 3.x Development Essentials - Dart Edition</i>. Payload Media, Inc."
    ]

    for p_item in pustaka_items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        run = p.add_run(p_item)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)

    doc.save(target_path)
    print("Dokumen Word berhasil dibuat di:", target_path)

if __name__ == "__main__":
    build_laporan_docx()
